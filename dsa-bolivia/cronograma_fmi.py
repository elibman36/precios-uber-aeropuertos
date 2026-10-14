"""Genera Cronograma_FMI_Bolivia.xlsx: desembolsos, repagos y cargos del SAF 2026-2029.

Uso: python cronograma_fmi.py [salida.xlsx]
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "Cronograma_FMI_Bolivia.xlsx"

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, italic=True, size=8)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
TOT_FILL = PatternFill("solid", fgColor="F2F2F2")
NUM = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;-0.0%;"-"'
PCT2 = '0.00%;-0.00%;"-"'
DATE = "mmm-yyyy"

wb = Workbook()
ws = wb.active
ws.title = "Supuestos"
wm = wb.create_sheet("Mensual")
wa = wb.create_sheet("Anual")
wn = wb.create_sheet("Notas_Fuentes")


def hdr(sheet, row, labels, col0=1):
    for j, lab in enumerate(labels):
        c = sheet.cell(row=row, column=col0 + j, value=lab)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


# ---------------------------------------------------------------------------
# SUPUESTOS
# ---------------------------------------------------------------------------
ws["A1"] = "Bolivia: Servicio Ampliado del FMI (SAF) 2026-2029. Supuestos y tramos"
ws["A1"].font = TITLE
ws["A2"] = ("Azul = insumo editable · Negro = fórmula · Amarillo = supuesto clave. "
            "Montos en millones de DEG salvo indicación.")
ws["A2"].font = NOTE

PARAMS = [
    ("quota", "Cuota de Bolivia en el FMI", 240.1, NUM, "DEG mm. Ficha del país, FMI"),
    ("access", "Acceso total aprobado", 1369.0, NUM, "DEG mm = 570% de la cuota. Directorio, 2-oct-2026"),
    ("usd_sdr", "Tipo de cambio USD por DEG", 211.3 / 156.1, '0.0000', "Implícito en el primer desembolso: USD 211,3 mm = DEG 156,1 mm (BCB, 6-oct-2026)"),
    ("sdr_rate", "Tasa de interés del DEG", 0.0287, PCT2, "Supuesto: implícita en la tasa de 3,47% informada, menos el margen de 60 pb. Se mantiene constante"),
    ("margin", "Margen sobre la tasa del DEG", 0.0060, PCT2, "60 pb desde la reforma de nov-2024"),
    ("lvl_thr", "Umbral de sobrecargo por nivel", 3.00, PCT, "300% de la cuota (reforma 2024)"),
    ("lvl_rate", "Sobrecargo por nivel", 0.0200, PCT2, "200 pb sobre el crédito que excede el umbral"),
    ("time_m", "Meses sobre el umbral para el sobrecargo por tiempo (SAF)", 51, '0', "51 meses en el SAF (36 en otros servicios)"),
    ("time_rate", "Sobrecargo por tiempo", 0.0075, PCT2, "75 pb adicionales (reforma 2024)"),
    ("svc", "Cargo por servicio (sobre cada desembolso)", 0.0050, PCT2, "50 pb por giro"),
    ("rep_start", "Inicio de repago tras cada desembolso", 54, '0" meses"', "4,5 años"),
    ("rep_end", "Último repago tras cada desembolso", 120, '0" meses"', "10 años"),
    ("rep_n", "Número de cuotas semestrales iguales", 12, '0', "Repago estándar del SAF"),
]
P = {}
hdr(ws, 4, ["Parámetro", "", "Valor", "Nota / fuente"])
for i, (k, lab, v, fmt, note) in enumerate(PARAMS):
    r = 5 + i
    P[k] = f"Supuestos!$C${r}"
    ws.cell(row=r, column=1, value=lab).font = BLACK
    c = ws.cell(row=r, column=3, value=v)
    c.font = BLUE
    c.number_format = fmt
    if k in ("sdr_rate", "usd_sdr"):
        c.fill = KEY_FILL
    ws.cell(row=r, column=4, value=note).font = NOTE
r_thr_amt = 5 + len(PARAMS)
ws.cell(row=r_thr_amt, column=1, value="Umbral de sobrecargo en DEG").font = BLACK
c = ws.cell(row=r_thr_amt, column=3, value=f"={P['lvl_thr']}*{P['quota']}")
c.number_format = NUM
P["thr_amt"] = f"Supuestos!$C${r_thr_amt}"
r_basic = r_thr_amt + 1
ws.cell(row=r_basic, column=1, value="Tasa de cargo básica (DEG + margen)").font = BLACK
c = ws.cell(row=r_basic, column=3, value=f"={P['sdr_rate']}+{P['margin']}")
c.number_format = PCT2
P["basic"] = f"Supuestos!$C${r_basic}"

# Tramos
TR0 = r_basic + 3
ws.cell(row=TR0 - 1, column=1, value="Tramos del acuerdo (fechas y montos editables)").font = BOLD
hdr(ws, TR0, ["#", "Evento", "Fecha", "Monto (DEG mm)", "Monto (USD mm)", "% de la cuota",
              "Acumulado (DEG mm)", "Acumulado (% cuota)", "Índice de mes",
              "Primer repago", "Último repago", "Cuota semestral (DEG mm)"])
ws.row_dimensions[TR0].height = 32
first_amt, first2 = 156.1, 156.0
rest = 1369.0 - first_amt - first2
tranches = [("Aprobación del Directorio", date(2026, 10, 1), first_amt),
            ("1ª revisión", date(2026, 12, 1), first2)]
rev_dates = [date(2027, 3, 1), date(2027, 6, 1), date(2027, 9, 1), date(2027, 12, 1),
             date(2028, 3, 1), date(2028, 6, 1), date(2028, 9, 1), date(2028, 12, 1),
             date(2029, 3, 1), date(2029, 6, 1), date(2029, 9, 1)]
for i, d in enumerate(rev_dates):
    tranches.append((f"{i + 2}ª revisión", d, round(rest / len(rev_dates), 2)))
NTR = len(tranches)
T1, TN = TR0 + 1, TR0 + NTR
for i, (ev, d, amt) in enumerate(tranches):
    r = TR0 + 1 + i
    ws.cell(row=r, column=1, value=i + 1).font = BLACK
    ws.cell(row=r, column=2, value=ev).font = BLACK
    c = ws.cell(row=r, column=3, value=d)
    c.font = BLUE
    c.number_format = DATE
    if i == NTR - 1:
        c = ws.cell(row=r, column=4, value=f"={P['access']}-SUM(D{T1}:D{r - 1})")
        c.font = BLACK
    else:
        c = ws.cell(row=r, column=4, value=amt)
        c.font = BLUE
        if i >= 1:
            c.fill = KEY_FILL  # monto supuesto
    c.number_format = NUM
    ws.cell(row=r, column=5, value=f"=D{r}*{P['usd_sdr']}").number_format = NUM
    ws.cell(row=r, column=6, value=f"=D{r}/{P['quota']}").number_format = PCT
    ws.cell(row=r, column=7, value=f"=SUM(D${T1}:D{r})").number_format = NUM
    ws.cell(row=r, column=8, value=f"=G{r}/{P['quota']}").number_format = PCT
    ws.cell(row=r, column=9, value=f"=YEAR(C{r})*12+MONTH(C{r})").number_format = '0'
    c = ws.cell(row=r, column=10, value=f"=EDATE(C{r},{P['rep_start']})")
    c.number_format = DATE
    c = ws.cell(row=r, column=11, value=f"=EDATE(C{r},{P['rep_end']})")
    c.number_format = DATE
    ws.cell(row=r, column=12, value=f"=D{r}/{P['rep_n']}").number_format = '#,##0.00'
    for cc in range(5, 13):
        ws.cell(row=r, column=cc).font = BLACK
rt = TN + 1
ws.cell(row=rt, column=2, value="Total").font = BOLD
for cc, f in [(4, f"=SUM(D{T1}:D{TN})"), (5, f"=SUM(E{T1}:E{TN})"), (6, f"=SUM(F{T1}:F{TN})")]:
    c = ws.cell(row=rt, column=cc, value=f)
    c.font = BOLD
    c.number_format = PCT if cc == 6 else NUM
    c.fill = TOT_FILL
ws.cell(row=rt + 1, column=2, value="Control: total = acceso").font = NOTE
c = ws.cell(row=rt + 1, column=4, value=f'=IF(ABS(D{rt}-{P["access"]})<0.001,"OK","REVISAR")')
c.font = BOLD
ws.cell(row=rt + 3, column=1, value=(
    "Fechas y montos de las revisiones 1 a 12: supuesto. El FMI no publicó el calendario en las fuentes consultadas. "
    "Los supuestos: (i) el primer tramo de DEG 156,1 mm es dato (BCB); (ii) un segundo tramo de monto similar en "
    "nov-dic 2026, según el ministro de Economía; (iii) desde 2027, revisiones trimestrales con montos iguales que "
    "completan DEG 1.369 mm, en un acuerdo de 36 meses (oct-2026 a sep-2029). El último tramo es residual, así que "
    "cualquier cambio en los montos mantiene el total.")).font = NOTE
ws.cell(row=rt + 3, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=rt + 3, start_column=1, end_row=rt + 3, end_column=12)
ws.row_dimensions[rt + 3].height = 54
for k, w in {"A": 46, "B": 26, "C": 12, "D": 14, "E": 14, "F": 11, "G": 14, "H": 13, "I": 10,
             "J": 12, "K": 12, "L": 14}.items():
    ws.column_dimensions[k].width = w
ws.column_dimensions["D"].width = 16

RNG_IDX = f"Supuestos!$I${T1}:$I${TN}"
RNG_AMT = f"Supuestos!$D${T1}:$D${TN}"

# ---------------------------------------------------------------------------
# MENSUAL
# ---------------------------------------------------------------------------
wm["A1"] = "Cronograma mensual (DEG mm)"
wm["A1"].font = TITLE
wm["A2"] = ("Los cargos se devengan sobre el saldo a fin del mes anterior; el FMI los liquida trimestralmente. "
            "Repago: 12 cuotas semestrales iguales de cada tramo, entre los meses 54 y 120 posteriores al giro.")
wm["A2"].font = NOTE
cols = ["Mes", "Índice", "Año", "Desembolso", "Repago de capital", "Saldo adeudado", "Saldo (% cuota)",
        "Exceso sobre umbral (saldo previo)", "Meses consecutivos sobre umbral", "Cargo básico",
        "Sobrecargo por nivel", "Sobrecargo por tiempo", "Cargo por servicio", "Cargos totales",
        "Servicio total (capital + cargos)"]
hdr(wm, 4, cols)
wm.row_dimensions[4].height = 44
start = date(2026, 10, 1)
n_months = (2039 - 2026) * 12 + (12 - 10) + 1  # oct-2026 .. dic-2039
M1 = 5
MN = M1 + n_months - 1
for i in range(n_months):
    r = M1 + i
    if i == 0:
        c = wm.cell(row=r, column=1, value=f"=Supuestos!$C${T1}")
        c.font = GREEN
    else:
        c = wm.cell(row=r, column=1, value=f"=EDATE(A{r - 1},1)")
    c.number_format = DATE
    wm.cell(row=r, column=2, value=f"=YEAR(A{r})*12+MONTH(A{r})")
    wm.cell(row=r, column=3, value=f"=YEAR(A{r})")
    wm.cell(row=r, column=4, value=f"=SUMIF({RNG_IDX},B{r},{RNG_AMT})")
    diff = f"(B{r}-{RNG_IDX})"
    wm.cell(row=r, column=5, value=(
        f"=SUMPRODUCT({RNG_AMT}*({diff}>={P['rep_start']})*({diff}<={P['rep_end']})"
        f"*(MOD({diff}-{P['rep_start']},6)=0))/{P['rep_n']}"))
    prevF = "0" if i == 0 else f"F{r - 1}"
    prevI = "0" if i == 0 else f"I{r - 1}"
    wm.cell(row=r, column=6, value=f"={prevF}+D{r}-E{r}")
    wm.cell(row=r, column=7, value=f"=F{r}/{P['quota']}")
    wm.cell(row=r, column=8, value=f"=MAX(0,{prevF}-{P['thr_amt']})")
    wm.cell(row=r, column=9, value=f"=IF(F{r}>{P['thr_amt']},{prevI}+1,0)")
    wm.cell(row=r, column=10, value=f"={prevF}*{P['basic']}/12")
    wm.cell(row=r, column=11, value=f"=H{r}*{P['lvl_rate']}/12")
    wm.cell(row=r, column=12, value=f"=IF({prevI}>{P['time_m']},H{r}*{P['time_rate']}/12,0)")
    wm.cell(row=r, column=13, value=f"=D{r}*{P['svc']}")
    wm.cell(row=r, column=14, value=f"=SUM(J{r}:M{r})")
    wm.cell(row=r, column=15, value=f"=E{r}+N{r}")
    for cc in range(2, 16):
        cell = wm.cell(row=r, column=cc)
        cell.font = BLACK
        cell.number_format = {2: '0', 3: '0', 7: PCT, 9: '0'}.get(cc, '#,##0.00;(#,##0.00);"-"')
# resaltar meses con giro o repago
wm.conditional_formatting.add(f"D{M1}:D{MN}", FormulaRule(formula=[f"D{M1}>0"], fill=PatternFill("solid", fgColor="DDEBF7")))
wm.conditional_formatting.add(f"E{M1}:E{MN}", FormulaRule(formula=[f"E{M1}>0"], fill=PatternFill("solid", fgColor="FCE4D6")))
wm.freeze_panes = "B5"
wm.column_dimensions["A"].width = 11
for cc in range(2, 16):
    wm.column_dimensions[get_column_letter(cc)].width = 13
wm.column_dimensions["B"].width = 8
wm.column_dimensions["C"].width = 7

# ---------------------------------------------------------------------------
# ANUAL
# ---------------------------------------------------------------------------
wa["A1"] = "Resumen anual: flujos con el FMI"
wa["A1"].font = TITLE
wa["A2"] = "DEG mm y USD mm (convertidos con el tipo de cambio USD/DEG de Supuestos, constante)."
wa["A2"].font = NOTE
acols = ["Año", "Desembolsos", "Repagos de capital", "Cargo básico", "Sobrecargos", "Cargo por servicio",
         "Cargos totales", "Servicio total", "Flujo neto (desemb. − servicio)", "Saldo a fin de año",
         "Saldo (% cuota)", "Desembolsos (USD)", "Repagos (USD)", "Cargos (USD)", "Servicio total (USD)",
         "Flujo neto (USD)", "Saldo fin de año (USD)"]
wa.cell(row=4, column=2, value="Millones de DEG").font = BOLD
wa.cell(row=4, column=12, value="Millones de USD").font = BOLD
hdr(wa, 5, acols)
wa.row_dimensions[5].height = 44
years = list(range(2026, 2040))
A1r = 6
ANr = A1r + len(years) - 1
MR = lambda c: f"Mensual!${c}${M1}:${c}${MN}"
for i, y in enumerate(years):
    r = A1r + i
    wa.cell(row=r, column=1, value=str(y)).font = BLACK
    yr = f"VALUE(A{r})"
    s = lambda c: f"=SUMIF({MR('C')},{yr},{MR(c)})"
    wa.cell(row=r, column=2, value=s("D"))
    wa.cell(row=r, column=3, value=s("E"))
    wa.cell(row=r, column=4, value=s("J"))
    wa.cell(row=r, column=5, value=f"=SUMIF({MR('C')},{yr},{MR('K')})+SUMIF({MR('C')},{yr},{MR('L')})")
    wa.cell(row=r, column=6, value=s("M"))
    wa.cell(row=r, column=7, value=f"=D{r}+E{r}+F{r}")
    wa.cell(row=r, column=8, value=f"=C{r}+G{r}")
    wa.cell(row=r, column=9, value=f"=B{r}-H{r}")
    wa.cell(row=r, column=10, value=f"=SUMIF({MR('B')},{yr}*12+12,{MR('F')})")
    wa.cell(row=r, column=11, value=f"=J{r}/{P['quota']}")
    for cc, src in [(12, "B"), (13, "C"), (14, "G"), (15, "H"), (16, "I"), (17, "J")]:
        wa.cell(row=r, column=cc, value=f"={src}{r}*{P['usd_sdr']}")
    for cc in range(2, 18):
        cell = wa.cell(row=r, column=cc)
        cell.font = BLACK
        cell.number_format = PCT if cc == 11 else NUM
rT = ANr + 1
wa.cell(row=rT, column=1, value="Total").font = BOLD
for cc in list(range(2, 10)) + list(range(12, 17)):
    L = get_column_letter(cc)
    c = wa.cell(row=rT, column=cc, value=f"=SUM({L}{A1r}:{L}{ANr})")
    c.font = BOLD
    c.number_format = NUM
    c.fill = TOT_FILL
wa.cell(row=rT, column=1).fill = TOT_FILL

r = rT + 2
wa.cell(row=r, column=1, value="Indicadores").font = BOLD
ind = [
    ("Saldo máximo (DEG mm)", f"=MAX({MR('F')})", NUM),
    ("Saldo máximo (% cuota)", f"=MAX({MR('G')})", PCT),
    ("Mes del saldo máximo", f"=INDEX({MR('A')},MATCH(MAX({MR('F')}),{MR('F')},0))", DATE),
    ("Primer repago de capital", f"=MIN(Supuestos!$J${T1}:$J${TN})", DATE),
    ("Último repago de capital", f"=MAX(Supuestos!$K${T1}:$K${TN})", DATE),
    ("Año de mayor servicio total (USD mm)", f"=INDEX(A{A1r}:A{ANr},MATCH(MAX(O{A1r}:O{ANr}),O{A1r}:O{ANr},0))", "@"),
    ("Servicio total en ese año (USD mm)", f"=MAX(O{A1r}:O{ANr})", NUM),
    ("Costo financiero total / monto girado", f"=G{rT}/B{rT}", PCT),
    ("Meses con sobrecargo por tiempo", f"=COUNTIF({MR('L')},\">0\")", '0'),
]
for j, (lab, f, fmt) in enumerate(ind):
    wa.cell(row=r + 1 + j, column=1, value=lab).font = BLACK
    c = wa.cell(row=r + 1 + j, column=3, value=f)
    c.font = BLACK
    c.number_format = fmt
wa.column_dimensions["A"].width = 36
for cc in range(2, 18):
    wa.column_dimensions[get_column_letter(cc)].width = 12
wa.freeze_panes = "B6"

# Gráficos
ch = BarChart()
ch.type = "col"
ch.grouping = "clustered"
ch.title = "Desembolsos vs. servicio total al FMI (USD mm)"
ch.y_axis.title = "USD mm"
ch.y_axis.number_format = "#,##0"
for src_col, color in [(12, "2A78D6"), (15, "EB6834")]:
    ref = Reference(wa, min_col=src_col, min_row=5, max_row=ANr)
    ch.add_data(ref, titles_from_data=True)
    ch.series[-1].graphicalProperties.solidFill = color
    ch.series[-1].graphicalProperties.line.solidFill = color
ch.set_categories(Reference(wa, min_col=1, min_row=A1r, max_row=ANr))
ch.legend.position = "b"
ch.gapWidth = 60
ch.height, ch.width = 9, 20
ch.x_axis.delete = False
ch.y_axis.delete = False
wa.add_chart(ch, f"A{r + len(ind) + 3}")

lc = LineChart()
lc.title = "Saldo adeudado al FMI (% de la cuota)"
lc.y_axis.number_format = "0%"
lc.add_data(Reference(wa, min_col=11, min_row=5, max_row=ANr), titles_from_data=True)
lc.series[0].graphicalProperties.line.solidFill = "2A78D6"
lc.series[0].graphicalProperties.line.width = 22000
lc.series[0].smooth = False
lc.set_categories(Reference(wa, min_col=1, min_row=A1r, max_row=ANr))
lc.legend = None
lc.height, lc.width = 9, 20
lc.x_axis.delete = False
lc.y_axis.delete = False
wa.add_chart(lc, f"J{r + len(ind) + 3}")

# ---------------------------------------------------------------------------
# NOTAS Y FUENTES
# ---------------------------------------------------------------------------
wn["A1"] = "Notas metodológicas y fuentes"
wn["A1"].font = TITLE
notes = [
    "Dato: el acuerdo SAF de 36 meses, por DEG 1.369 mm (570% de la cuota), fue aprobado el 2-oct-2026. El primer giro fue de DEG 156,1 mm (USD 211,3 mm), recibido por el BCB.",
    "Supuesto: el calendario de revisiones posteriores. Las fuentes consultadas no publican el cuadro de fases del staff report. Se puede editar en Supuestos (fechas y montos).",
    "Repago: cada giro se devuelve en 12 cuotas semestrales iguales, desde los 4,5 años hasta los 10 años posteriores al giro.",
    "Cargo básico: tasa de interés del DEG + 60 pb. La tasa del DEG es variable (se fija semanalmente); acá se supone constante en 2,87%, el valor implícito en el 3,47% informado.",
    "Sobrecargos: 200 pb sobre el crédito que excede el 300% de la cuota. Se suman 75 pb si ese exceso se mantiene más de 51 meses (SAF).",
    "Cargo por servicio: 50 pb sobre cada giro. No se incluye la comisión de compromiso porque se reembolsa a medida que se gira; sólo sería un costo neto si no se completa el acceso.",
    "Simplificaciones: los cargos se devengan mensualmente sobre el saldo del mes anterior (el FMI liquida trimestralmente y por saldo diario), y la conversión a USD usa un tipo USD/DEG constante.",
    "Esta hoja reemplaza la aproximación anual del FMI usada en DSA_Bolivia.xlsx (USD 0,43 / 0,62 / 0,62 / 0,23 mm M en 2026-29, con repagos lineales de 1/6 por año).",
]
for i, t in enumerate(notes):
    c = wn.cell(row=3 + i, column=1, value=f"• {t}")
    c.font = BLACK
    c.alignment = Alignment(wrap_text=True, vertical="top")
    wn.row_dimensions[3 + i].height = 30
r = 4 + len(notes)
wn.cell(row=r, column=1, value="Fuentes").font = BOLD
SRC = [
    ("FMI, comunicado de aprobación del SAF (2-oct-2026)", "https://www.imf.org/es/news/articles/2026/10/02/pr26317-bolivia-imf-approves-36-month-eff-arrangement"),
    ("BCB confirma el primer desembolso: USD 211,3 mm = DEG 156,1 mm", "https://eldeber.com.bo/economia/banco-central-bolivia-confirma-primer-desembolso-us-211-3-millones-fmi_1791317879"),
    ("Condiciones: tasa 3,47%, repago desde 4 años y 6 meses hasta 10 años", "https://www.economy.com.bo/articulo/economia/fmi-aprueba-financiamiento-us-1900-millones-bolivia-desembolso-inmediato-us-214-millones/20261002185448024301.html"),
    ("Ministro: segundo desembolso nov-dic 2026, luego trimestrales", "https://www.reduno.com.bo/economia/bolivia-recibira-us-214-millones-del-fmi-para-fortalecer-reservas-y-preve-nuevos-desembolsos-202610674812"),
    ("FMI, reforma de cargos y sobrecargos (oct-2024)", "https://www.imf.org/en/News/Articles/2024/10/21/pr-24385-imf-concludes-the-review-of-charges-and-surcharge-policy-and-approves-reforms"),
    ("FMI, preguntas frecuentes sobre cargos y sobrecargos", "https://www.imf.org/en/about/faq/charges-and-surcharge-policy"),
]
for i, (a, b) in enumerate(SRC):
    wn.cell(row=r + 1 + i, column=1, value=a).font = BLACK
    c = wn.cell(row=r + 1 + i, column=2, value=b)
    c.hyperlink = b
    c.font = Font(name=FONT, color="0563C1", underline="single", size=9)
wn.column_dimensions["A"].width = 110
wn.column_dimensions["B"].width = 100

wb.save(OUT)
print("ok", OUT, "meses", n_months, "tramos", NTR)
