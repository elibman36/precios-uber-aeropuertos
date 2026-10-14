"""Genera DSA_Bolivia.xlsx: el modelo de modelo.py escrito como fórmulas vivas,
más resultados de escenarios y Monte Carlo (valores calculados por modelo.py).

Uso: python construir_excel.py [salida.xlsx] [escenario_activo]
"""
import sys
import numpy as np
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import modelo as m

OUT = sys.argv[1] if len(sys.argv) > 1 else "DSA_Bolivia.xlsx"
ACTIVE = int(sys.argv[2]) if len(sys.argv) > 2 else 1

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
EST_FILL = PatternFill("solid", fgColor="F2F2F2")
RED_FILL = PatternFill("solid", fgColor="F8CBAD")
GRN_FILL = PatternFill("solid", fgColor="C6EFCE")
THIN = Side(style="thin", color="BFBFBF")

FMT = {
    "pct": '0.0%;-0.0%;"-"',
    "usd": '0.00;(0.00);"-"',
    "bs": '#,##0.0;(#,##0.0);"-"',
    "idx": '0.0',
    "er": '0.00',
    "x": '0.00"x"',
    "m": '0.0" m"',
}

YEARS = m.YEARS
T = m.T
FIRST_COL = 3


def col(t):
    return get_column_letter(FIRST_COL + t)


wb = Workbook()
ws_port = wb.active
ws_port.title = "Portada"
ws_sup = wb.create_sheet("Supuestos")
ws_ext = wb.create_sheet("Externo")
ws_fis = wb.create_sheet("Fiscal_Deuda")
ws_ind = wb.create_sheet("Indicadores")
ws_esc = wb.create_sheet("Escenarios")
ws_mc = wb.create_sheet("Estocastico")
ws_cmp = wb.create_sheet("FMI_vs_Modelo")
ws_src = wb.create_sheet("Fuentes")

SHEET_NAME = {"sup": "Supuestos", "ext": "Externo", "fis": "Fiscal_Deuda", "ind": "Indicadores"}
ROWS = {"sup": {}, "ext": {}, "fis": {}, "ind": {}}
SCAL_ROW = {}
PARAM_COL = {}
PARAM_ACTIVE_ROW = None


def setw(ws, widths):
    for k, v in widths.items():
        ws.column_dimensions[k].width = v


def year_header(ws, row):
    ws.cell(row=row, column=1, value="Concepto").font = HDR_FONT
    ws.cell(row=row, column=2, value="Unidad").font = HDR_FONT
    for t, y in enumerate(YEARS):
        c = ws.cell(row=row, column=FIRST_COL + t, value=str(y))
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center")
    for cc in range(1, FIRST_COL + T + 1):
        ws.cell(row=row, column=cc).fill = HDR_FILL
    ws.cell(row=row, column=FIRST_COL + T, value="Nota / fuente").font = HDR_FONT
    # sombreado de columnas 2025 (dato) y 2026 (estimación)
    for t in (0, 1):
        ws.cell(row=row - 1, column=FIRST_COL + t, value="dato" if t == 0 else "estim.").font = Font(name=FONT, italic=True, size=8)


def ref(sheet, key, t, cur=None, absolute=False):
    r = ROWS[sheet][key]
    c = col(t)
    a = f"${c}${r}" if absolute else f"{c}{r}"
    return a if sheet == cur else f"{SHEET_NAME[sheet]}!{a}"


def SC(key):
    return f"Supuestos!$C${SCAL_ROW[key]}"


def P(key):
    return f"Supuestos!${PARAM_COL[key]}${PARAM_ACTIVE_ROW}"


def section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = BOLD
    for cc in range(1, FIRST_COL + T + 1):
        ws.cell(row=row, column=cc).fill = SEC_FILL


def is_link(v):
    if not isinstance(v, str) or not v.startswith("="):
        return False
    body = v[1:]
    return "!" in body and not any(op in body for op in "+-*/(),")


def write_row(ws, sheet, row, key, label, unit, fn, fmt, note=None, bold=False, key_fill=False):
    ws.cell(row=row, column=1, value=label).font = BOLD if bold else BLACK
    ws.cell(row=row, column=2, value=unit).font = Font(name=FONT, italic=True, size=9)
    for t in range(T):
        v = fn(t)
        if v is None:
            continue
        c = ws.cell(row=row, column=FIRST_COL + t, value=v)
        if isinstance(v, str) and v.startswith("="):
            c.font = GREEN if is_link(v) else BLACK
        else:
            c.font = BLUE
        if bold:
            c.font = Font(name=FONT, bold=True, color=c.font.color.rgb if c.font.color else "000000")
        c.number_format = FMT[fmt]
        if key_fill:
            c.fill = KEY_FILL
        elif t <= 1:
            c.fill = EST_FILL
    if note:
        ws.cell(row=row, column=FIRST_COL + T, value=note).font = Font(name=FONT, size=8, italic=True)


# ---------------------------------------------------------------------------
# SUPUESTOS
# ---------------------------------------------------------------------------
SUP_SERIES = [
    ("Macro", None),
    ("g_base", ("Crecimiento real (base)", "%", "pct", "2025: INE. 2026-28: FMI (WEO oct-2026 / staff report EFF). 2029+: supuesto propio")),
    ("defl_base", ("Deflactor del PIB (base)", "%", "pct", "Supuesto propio; FMI: IPC prom. 12,7% (2026), 15,8% (2027)")),
    ("pi_us", ("Inflación EE.UU.", "%", "pct", "Supuesto")),
    ("rer_base", ("Var. tipo de cambio real (base; + = depreciación)", "%", "pct", "Corrección parcial del overshooting post-flotación (jun-2026)")),
    ("er_input", ("Tipo de cambio fin de año (dato/estim.)", "Bs/USD", "er", "6,96 (2025, oficial); ~11,7-11,8 oct-2026 (BCB, flotación desde RM 245/2026)")),
    ("Precios externos (índice 2025=100)", None),
    ("p_prec", ("Precios metales preciosos (plata/oro, export.)", "índice", "idx", "Calibrado: export. minerales +91% en valor con -12% volumen (IBCE, 1S-2026)")),
    ("p_basem", ("Precios metales base (zinc, estaño, plomo)", "índice", "idx", "Supuesto")),
    ("p_oil", ("Precio petróleo (Brent)", "índice", "idx", "Supuesto; gas a Brasil indexado con 1 año de rezago")),
    ("p_gold", ("Precio del oro (valuación reservas)", "índice", "idx", "Supuesto")),
    ("p_agro_g", ("Var. precios agro (soya y otros)", "%", "pct", "Supuesto: plano")),
    ("Volúmenes (var. anual)", None),
    ("v_prec", ("Volumen metales preciosos", "%", "pct", "Supuesto")),
    ("v_basem", ("Volumen metales base", "%", "pct", "Supuesto")),
    ("v_hydro", ("Volumen exportado de gas", "%", "pct", "Declinación de campos maduros; export. de gas -33% en 2025 (INE)")),
    ("v_agro", ("Volumen agro", "%", "pct", "Supuesto")),
    ("v_fuel", ("Volumen importado de combustibles", "%", "pct", "Fin de subsidios reduce contrabando; luego crece con actividad")),
    ("Balanza de pagos: valores iniciales (USD miles de millones)", None),
    ("x_prec0", ("Export. metales preciosos", "USD mm M", "usd", "2025: BCB/INE aprox.; 2026: anualización ene-ago (INE: USD 8.987 M acumulado)")),
    ("x_basem0", ("Export. metales base", "USD mm M", "usd", "Ídem")),
    ("x_hydro0", ("Export. hidrocarburos", "USD mm M", "usd", "2025: gas USD 1.077 M + otros (INE); 1S-2026: USD 498 M")),
    ("x_agro0", ("Export. agro (soya y derivados, otros)", "USD mm M", "usd", "Ídem")),
    ("x_other0", ("Export. resto (manufacturas no tradicionales)", "USD mm M", "usd", "Ídem")),
    ("m_fuel0", ("Import. combustibles", "USD mm M", "usd", "Estimación propia")),
    ("m_nf0", ("Import. resto (FOB)", "USD mm M", "usd", "2025 calibrado a balanza de bienes +0,33 (BCB)")),
    ("serv0", ("Servicios (neto)", "USD mm M", "usd", "BCB, BoP 2025; 2026 estim.")),
    ("sec0", ("Ingreso secundario (remesas, neto)", "USD mm M", "usd", "BCB; 2026 estim.")),
    ("Cuenta financiera (USD miles de millones, flujos)", None),
    ("fdi", ("IED neta", "USD mm M", "usd", "Supuesto: recuperación gradual con nueva ley de hidrocarburos/litio")),
    ("disb", ("Desembolsos deuda externa pública (excl. FMI)", "USD mm M", "usd", ">USD 1.000 M BM/BID nov-dic 2026; ~USD 5.000 M esperados de otros acreedores durante el programa")),
    ("amort", ("Amortizaciones deuda externa pública", "USD mm M", "usd", "Estimación propia de cronograma (multilateral + bilateral + bonos 2028/2030 en manos externas)")),
    ("reprof_relief", ("Reperfilamiento: alivio (bilateral + bonos)", "USD mm M", "usd", "Sólo activo en escenario de reperfilamiento: diferimiento 5 años")),
    ("reprof_repay", ("Reperfilamiento: repago diferido", "USD mm M", "usd", "Ídem")),
    ("imf_disb", ("Desembolsos FMI (SAF)", "USD mm M", "usd", "SAF 36 meses, DEG 1.369 M (~USD 1.900 M); USD 214 M en oct-2026")),
    ("priv", ("Otros flujos privados netos (incl. errores y omisiones)", "USD mm M", "usd", "2026 calibrado a reservas ~5,2 mm M fin de año")),
    ("Fiscal y deuda", None),
    ("pb_nc", ("Balance primario no-commodity (SPNF)", "% PIB", "pct", "Calibrado para replicar el déficit global del programa: 9,2% (2026), 6,7% (2027), 4,5% (2028), 3,2% (2029), 3,0% después")),
    ("sub_share", ("Subsidio a combustibles / import. combustibles", "%", "pct", "Eliminación dic-2025; remanente se elimina en 2027 (mecanismo automático)")),
    ("i_dom", ("Tasa efectiva deuda interna", "%", "pct", "Supuesto: sube a medida que se corta financiamiento BCB")),
    ("i_ext", ("Tasa efectiva deuda externa pública", "%", "pct", "Supuesto (multilateral SOFR+, bonos 4,5%/7,5%)")),
    ("bcb_share", ("Deuda interna en moneda local en manos del BCB", "%", "pct", "BCB ~54,6% de la deuda interna (Jubileo/ANF); cae 2 pp/año")),
    ("ddom_usd", ("Deuda interna en USD o indexada (MVDOL)", "USD mm M", "usd", "Supuesto; incluye bonos canjeados a MVDOL (mar-2026)")),
]

ws_sup["A1"] = "Supuestos del modelo"
ws_sup["A1"].font = TITLE
ws_sup["A2"] = ("Celdas en azul = insumos editables. Gris = 2025 (dato) y 2026 (estimación). Amarillo = supuestos clave. "
                "Los escenarios aplican choques sobre estas sendas a partir de 2027.")
ws_sup["A2"].font = Font(name=FONT, italic=True, size=9)
year_header(ws_sup, 4)
r = 5
for key, meta in SUP_SERIES:
    if meta is None:
        section(ws_sup, r, key)
        r += 1
        continue
    ROWS["sup"][key] = r
    label, unit, fmt, note = meta
    vals = m.SERIES[key]

    def fn(t, vals=vals, key=key):
        v = vals[t]
        if v is None:
            return None
        if key == "pb_nc" and t == 0:
            return v
        return float(v)
    write_row(ws_sup, "sup", r, key, label, unit, fn, fmt, note,
              key_fill=key in ("p_prec", "v_hydro", "pb_nc", "rer_base"))
    r += 1

# Escalares
r += 1
section(ws_sup, r, "Parámetros (elasticidades, umbrales y anclas 2025)")
r += 1
SCAL_META = {
    "theta_h": ("Renta fiscal hidrocarburos / valor exportado", "pct", "IDH 32% + regalías 18% sobre valor en boca de pozo, neto"),
    "theta_m": ("Renta fiscal minera / valor exportado", "pct", "Regalías + IUE; baja por peso de cooperativas"),
    "eps_x": ("Elasticidad export. no tradicionales al TCR", "usd", "Supuesto"),
    "eps_m": ("Elasticidad import. al TCR", "usd", "Supuesto"),
    "eta_y": ("Elasticidad ingreso de importaciones", "usd", "Supuesto"),
    "eps_sr": ("Elasticidad de corto plazo para cerrar brecha externa", "usd", "Devaluación requerida = brecha / (elasticidad x import.)"),
    "kappa": ("Costo en crecimiento por punto de devaluación de ajuste", "usd", "Devaluación contractiva (efecto hoja de balance + compresión de import.)"),
    "pt": ("Pass-through de la devaluación al deflactor", "pct", "Supuesto"),
    "floor_m": ("Piso duro de reservas", "m", "Meses de importaciones"),
    "target_m": ("Nivel cómodo de reservas", "m", "Meses de importaciones"),
    "lam": ("Fracción de la brecha vs nivel cómodo cerrada vía tipo de cambio", "pct", "Regla de flotación administrada"),
    "mu": ("Propensión a importar del impulso fiscal", "usd", "Déficits gemelos: cada punto de desliz fiscal eleva import."),
    "pi_world": ("Inflación externa en USD", "pct", "Supuesto"),
    "w_nom": ("Crecimiento nominal de demanda externa", "pct", "Supuesto"),
    "prof_share": ("Remisión de utilidades / export. extractivas", "pct", "Fuga del boom minero vía ingreso primario"),
    "remit_g": ("Crecimiento de remesas", "pct", "Supuesto"),
    "imf_rate": ("Costo del crédito FMI (tasa DEG + margen + recargos)", "pct", "Supuesto"),
    "dom_mat": ("Vida media deuda interna", "usd", "Años; determina amortizaciones internas en NBF"),
    "st_share": ("Fracción de deuda externa privada que vence en el año", "pct", "Para ratio Greenspan-Guidotti"),
    "priv_ext0": ("Deuda externa privada 2026", "usd", "USD mm M; supuesto"),
    "y0": ("PIB nominal 2025", "bs", "Bs mm M = USD 55.000 M (CEPAL) x 6,96"),
    "dext0": ("Deuda externa pública dic-2025", "usd", "USD mm M; BCB (USD 14.117 M a sep-2026 = 47 M menos que dic-2025)"),
    "ddom_lc0": ("Deuda interna en moneda local dic-2025", "bs", "Bs mm M; calibrado para deuda total = 83,3% PIB (FMI)"),
    "gold0": ("Reservas en oro dic-2025", "usd", "USD mm M; estimación"),
    "ngold0": ("Reservas no-oro dic-2025", "usd", "USD mm M; total ~3.813 M (MEFP)"),
    "ob0": ("Balance global SPNF 2025", "pct", "FMI: -11,4% PIB"),
    "pb0": ("Balance primario SPNF 2025", "pct", "Estimación"),
    "imfout0": ("Crédito FMI vigente dic-2025", "usd", "USD mm M"),
}
for k, v in m.SCALARS.items():
    label, fmt, note = SCAL_META[k]
    SCAL_ROW[k] = r
    ws_sup.cell(row=r, column=1, value=label).font = BLACK
    c = ws_sup.cell(row=r, column=3, value=v)
    c.font = BLUE
    c.number_format = FMT[fmt] if fmt != "usd" else "0.00"
    ws_sup.cell(row=r, column=4, value=note).font = Font(name=FONT, size=8, italic=True)
    r += 1

# Tabla de escenarios
r += 1
section(ws_sup, r, "Escenarios: parámetros de choque (se aplican desde 2027)")
r += 1
PARAM_LABEL = {
    "sh_prec": "Shock nivel precios metales preciosos y oro",
    "sh_basem": "Shock nivel metales base",
    "sh_oil": "Shock nivel petróleo",
    "hydro_extra": "Declinación extra vol. gas (pp/año)",
    "fis_slip": "Desliz fiscal (pp PIB, bal. primario)",
    "g_shock": "Shock crecimiento 2027-28 (pp)",
    "rer_adj": "Ajuste TCR 2027-29 (pp/año)",
    "priv_shift": "Shock flujos privados (USD mm M/año)",
    "fdi_mult": "Multiplicador IED",
    "reprof": "Reperfilamiento (1 = sí)",
    "sub_return": "Retorno de subsidio (fracción import. comb.)",
    "cont": "Pasivos contingentes 2027 (% PIB)",
}
ws_sup.cell(row=r, column=1, value="#").font = HDR_FONT
ws_sup.cell(row=r, column=2, value="Escenario").font = HDR_FONT
for j, k in enumerate(m.PARAM_KEYS):
    cl = get_column_letter(3 + j)
    PARAM_COL[k] = cl
    c = ws_sup.cell(row=r, column=3 + j, value=PARAM_LABEL[k])
    c.font = HDR_FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")
for cc in range(1, 3 + len(m.PARAM_KEYS)):
    ws_sup.cell(row=r, column=cc).fill = HDR_FILL
ws_sup.row_dimensions[r].height = 54
r += 1
SCEN = m.scenario_params()
scen_first = r
for i, (name, q) in enumerate(SCEN):
    ws_sup.cell(row=r, column=1, value=i + 1).font = BLACK
    ws_sup.cell(row=r, column=2, value=name).font = BLACK
    for j, k in enumerate(m.PARAM_KEYS):
        c = ws_sup.cell(row=r, column=3 + j, value=float(q[k]))
        c.font = BLUE
        c.number_format = '0.0%;-0.0%;"-"' if k not in ("priv_shift", "fdi_mult", "reprof") else '0.00;-0.00;"-"'
    if name.startswith("Réplica"):
        ws_sup.cell(row=r, column=3 + m.PARAM_KEYS.index("rer_adj")).comment = Comment(
            "Calibrado en modelo.py: ajuste del TCR 2027-29 que lleva la deuda a 89,2% del PIB en 2029 "
            "(cifra del FMI) con los mismos déficits del programa.", "DSA")
    r += 1
scen_last = r - 1
PARAM_ACTIVE_ROW = r
ws_sup.cell(row=r, column=1, value="►").font = BOLD
ws_sup.cell(row=r, column=2, value="Escenario activo (según Portada)").font = BOLD
for j, k in enumerate(m.PARAM_KEYS):
    cl = get_column_letter(3 + j)
    c = ws_sup.cell(row=r, column=3 + j,
                    value=f"=INDEX({cl}${scen_first}:{cl}${scen_last},Portada!$C$6)")
    c.font = BLACK
    c.fill = KEY_FILL
    c.number_format = ws_sup.cell(row=scen_first, column=3 + j).number_format
setw(ws_sup, {"A": 52, "B": 16, **{col(t): 10 for t in range(T)}, col(T): 70})
ws_sup.freeze_panes = "C5"

# ---------------------------------------------------------------------------
# EXTERNO y FISCAL: especificación de filas
# ---------------------------------------------------------------------------
def S(key, t):
    return ref("sup", key, t)


def E(key, t):
    return ref("ext", key, t, cur=CUR[0])


def F(key, t):
    return ref("fis", key, t, cur=CUR[0])


def I(key, t):
    return ref("ind", key, t, cur=CUR[0])


CUR = [None]

EXT_SPEC = [
    ("Precios efectivos (escenario activo)", None),
    ("P_prec", "Precios metales preciosos", "índice", "idx",
     lambda t: f"={S('p_prec', t)}" if t < 2 else f"={S('p_prec', t)}*(1+{P('sh_prec')})"),
    ("P_basem", "Precios metales base", "índice", "idx",
     lambda t: f"={S('p_basem', t)}" if t < 2 else f"={S('p_basem', t)}*(1+{P('sh_basem')})"),
    ("P_oil", "Precio petróleo", "índice", "idx",
     lambda t: f"={S('p_oil', t)}" if t < 2 else f"={S('p_oil', t)}*(1+{P('sh_oil')})"),
    ("P_gold", "Precio oro", "índice", "idx",
     lambda t: f"={S('p_gold', t)}" if t < 2 else f"={S('p_gold', t)}*(1+{P('sh_prec')})"),
    ("rer", "Var. TCR exógena", "%", "pct",
     lambda t: None if t == 0 else (f"={S('rer_base', t)}+{P('rer_adj')}" if t in (2, 3, 4) else f"={S('rer_base', t)}")),
    ("rd", "Depreciación real efectiva (incl. rezago de ajuste)", "%", "pct",
     lambda t: None if t == 0 else f"={E('rer', t)}+(1-{SC('pt')})*{E('d', t - 1)}"),
    ("g_pre", "Crecimiento real antes del ajuste externo", "%", "pct",
     lambda t: f"={S('g_base', t)}" if t not in (2, 3) else f"={S('g_base', t)}+{P('g_shock')}"),
    ("Cuenta corriente (USD miles de millones)", None),
    ("X_prec", "Export. metales preciosos", "USD mm M", "usd",
     lambda t: f"={S('x_prec0', t)}" if t < 2 else f"={E('X_prec', t - 1)}*{E('P_prec', t)}/{E('P_prec', t - 1)}*(1+{S('v_prec', t)})"),
    ("X_basem", "Export. metales base", "USD mm M", "usd",
     lambda t: f"={S('x_basem0', t)}" if t < 2 else f"={E('X_basem', t - 1)}*{E('P_basem', t)}/{E('P_basem', t - 1)}*(1+{S('v_basem', t)})"),
    ("X_hydro", "Export. hidrocarburos", "USD mm M", "usd",
     lambda t: f"={S('x_hydro0', t)}" if t < 2 else f"={E('X_hydro', t - 1)}*{E('P_oil', t - 1)}/{E('P_oil', t - 2)}*(1+{S('v_hydro', t)}+{P('hydro_extra')})"),
    ("X_agro", "Export. agro", "USD mm M", "usd",
     lambda t: f"={S('x_agro0', t)}" if t < 2 else f"={E('X_agro', t - 1)}*(1+{S('p_agro_g', t)})*(1+{S('v_agro', t)})"),
    ("X_other", "Export. resto", "USD mm M", "usd",
     lambda t: f"={S('x_other0', t)}" if t < 2 else f"={E('X_other', t - 1)}*(1+{SC('w_nom')})*(1+{SC('eps_x')}*{E('rd', t)})"),
    ("X", "Exportaciones de bienes", "USD mm M", "usd",
     lambda t: f"=SUM({col(t)}{ROWS['ext']['X_prec']}:{col(t)}{ROWS['ext']['X_other']})"),
    ("M_fuel", "Import. combustibles", "USD mm M", "usd",
     lambda t: f"={S('m_fuel0', t)}" if t < 2 else f"={E('M_fuel', t - 1)}*{E('P_oil', t)}/{E('P_oil', t - 1)}*(1+{S('v_fuel', t)})"),
    ("M_nf_pre", "Import. resto antes del ajuste", "USD mm M", "usd",
     lambda t: f"={S('m_nf0', t)}" if t < 2 else
     (f"={E('M_nf', t - 1)}*(1+{SC('pi_world')})*(1+{SC('eta_y')}*{E('g_pre', t)})*(1-{SC('eps_m')}*{E('rd', t)})"
      f"-{SC('mu')}*{P('fis_slip')}*{F('Yusd', t - 1)}")),
    ("M_nf", "Import. resto (después del ajuste)", "USD mm M", "usd",
     lambda t: f"={E('M_nf_pre', t)}" if t == 0 else f"={E('M_nf_pre', t)}-{E('gap', t)}"),
    ("M", "Importaciones de bienes", "USD mm M", "usd",
     lambda t: f"={E('M_fuel', t)}+{E('M_nf', t)}"),
    ("serv", "Servicios (neto)", "USD mm M", "usd",
     lambda t: f"={S('serv0', t)}" if t < 2 else f"={E('serv', t - 1)}*(1+{SC('pi_world')})"),
    ("prof", "Utilidades remitidas (extractivas)", "USD mm M", "usd",
     lambda t: f"=-{SC('prof_share')}*({E('X_prec', t)}+{E('X_basem', t)}+{E('X_hydro', t)})"),
    ("pubint", "Intereses deuda externa pública", "USD mm M", "usd",
     lambda t: f"=-{S('i_ext', t)}*{F('Dext', 0)}" if t == 0 else f"=-{S('i_ext', t)}*{F('Dext', t - 1)}"),
    ("imfch", "Cargos FMI", "USD mm M", "usd",
     lambda t: 0.0 if t == 0 else f"=-{SC('imf_rate')}*{E('imf_out', t - 1)}"),
    ("sec", "Ingreso secundario (remesas)", "USD mm M", "usd",
     lambda t: f"={S('sec0', t)}" if t < 2 else f"={E('sec', t - 1)}*(1+{SC('remit_g')})"),
    ("CA_pre", "Cuenta corriente antes del ajuste", "USD mm M", "usd",
     lambda t: None if t == 0 else
     f"={E('X', t)}-{E('M_fuel', t)}-{E('M_nf_pre', t)}+{E('serv', t)}+{E('prof', t)}+{E('pubint', t)}+{E('imfch', t)}+{E('sec', t)}"),
    ("CA", "Cuenta corriente", "USD mm M", "usd",
     lambda t: f"={E('X', t)}-{E('M', t)}+{E('serv', t)}+{E('prof', t)}+{E('pubint', t)}+{E('imfch', t)}+{E('sec', t)}"
     if t == 0 else f"={E('CA_pre', t)}+{E('gap', t)}"),
    ("Cuenta financiera y reservas (USD miles de millones)", None),
    ("fdi", "IED neta", "USD mm M", "usd",
     lambda t: f"={S('fdi', t)}" if t < 2 else f"={S('fdi', t)}*{P('fdi_mult')}"),
    ("disb", "Desembolsos deuda externa pública", "USD mm M", "usd", lambda t: f"={S('disb', t)}"),
    ("amort", "Amortizaciones deuda externa pública", "USD mm M", "usd",
     lambda t: f"={S('amort', t)}" if t == 0 else
     f"={S('amort', t)}+{P('reprof')}*({S('reprof_repay', t)}-{S('reprof_relief', t)})"),
    ("imf_disb", "Desembolsos FMI", "USD mm M", "usd", lambda t: f"={S('imf_disb', t)}"),
    ("imf_rep", "Recompras FMI (4,5-10 años)", "USD mm M", "usd",
     lambda t: 0.0 if t == 0 else ("=(" + "+".join(S('imf_disb', t - k) for k in range(5, 11) if t - k >= 0) + ")/6"
                                   if t >= 5 else 0.0)),
    ("imf_out", "Crédito FMI vigente", "USD mm M", "usd",
     lambda t: f"={SC('imfout0')}" if t == 0 else f"={E('imf_out', t - 1)}+{E('imf_disb', t)}-{E('imf_rep', t)}"),
    ("priv", "Otros flujos privados netos", "USD mm M", "usd",
     lambda t: f"={S('priv', t)}" if t < 2 else f"={S('priv', t)}+{P('priv_shift')}"),
    ("gold", "Reservas en oro (valuación)", "USD mm M", "usd",
     lambda t: f"={SC('gold0')}" if t == 0 else f"={E('gold', t - 1)}*{E('P_gold', t)}/{E('P_gold', t - 1)}"),
    ("ngold_pre", "Reservas no-oro antes del ajuste", "USD mm M", "usd",
     lambda t: f"={SC('ngold0')}" if t == 0 else
     (f"={E('R', t - 1)}-{E('gold', t - 1)}+{E('CA_pre', t)}+{E('fdi', t)}+{E('disb', t)}-{E('amort', t)}"
      f"+{E('imf_disb', t)}-{E('imf_rep', t)}+{E('priv', t)}")),
    ("R_pre", "Reservas antes del ajuste", "USD mm M", "usd",
     lambda t: None if t == 0 else f"={E('ngold_pre', t)}+{E('gold', t)}"),
    ("floor", "Piso duro de reservas", "USD mm M", "usd",
     lambda t: None if t == 0 else f"={SC('floor_m')}/12*({E('M_fuel', t)}+{E('M_nf_pre', t)})"),
    ("target", "Nivel cómodo de reservas", "USD mm M", "usd",
     lambda t: None if t == 0 else f"={SC('target_m')}/12*({E('M_fuel', t)}+{E('M_nf_pre', t)})"),
    ("gap", "Brecha externa cerrada vía tipo de cambio", "USD mm M", "usd",
     lambda t: 0.0 if t < 2 else
     f"=MAX(0,{E('floor', t)}-{E('R_pre', t)},{SC('lam')}*({E('target', t)}-{E('R_pre', t)}))"),
    ("d", "Depreciación nominal de ajuste (endógena)", "%", "pct",
     lambda t: 0.0 if t == 0 else f"={E('gap', t)}/({SC('eps_sr')}*{E('M_nf_pre', t)})"),
    ("R", "Reservas internacionales brutas", "USD mm M", "usd",
     lambda t: f"={E('gold', t)}+{E('ngold_pre', t)}" if t == 0 else f"={E('R_pre', t)}+{E('gap', t)}"),
]

FIS_SPEC = [
    ("Macro", None),
    ("g", "Crecimiento real", "%", "pct",
     lambda t: f"={S('g_base', t)}" if t == 0 else f"={E('g_pre', t)}-{SC('kappa')}*{E('d', t)}"),
    ("pi", "Deflactor del PIB", "%", "pct",
     lambda t: f"={S('defl_base', t)}" if t == 0 else f"={S('defl_base', t)}+{SC('pt')}*{E('d', t)}"),
    ("Y", "PIB nominal", "Bs mm M", "bs",
     lambda t: f"={SC('y0')}" if t == 0 else f"={F('Y', t - 1)}*(1+{F('g', t)})*(1+{F('pi', t)})"),
    ("ER", "Tipo de cambio fin de año", "Bs/USD", "er",
     lambda t: f"={S('er_input', t)}" if t < 2 else
     f"={F('ER', t - 1)}*(1+{S('defl_base', t)})/(1+{S('pi_us', t)})*(1+{E('rer', t)})*(1+{E('d', t)})"),
    ("ERavg", "Tipo de cambio promedio", "Bs/USD", "er",
     lambda t: f"={F('ER', t)}" if t == 0 else f"=({F('ER', t - 1)}+{F('ER', t)})/2"),
    ("Yusd", "PIB en dólares", "USD mm M", "usd", lambda t: f"={F('Y', t)}/{F('ERavg', t)}"),
    ("Sector público no financiero (Bs miles de millones)", None),
    ("rev_h", "Renta hidrocarburífera (vinculada a export.)", "Bs mm M", "bs",
     lambda t: None if t == 0 else f"={SC('theta_h')}*{E('X_hydro', t)}*{F('ERavg', t)}"),
    ("rev_m", "Renta minera", "Bs mm M", "bs",
     lambda t: None if t == 0 else f"={SC('theta_m')}*({E('X_prec', t)}+{E('X_basem', t)})*{F('ERavg', t)}"),
    ("sub", "Subsidio a combustibles", "Bs mm M", "bs",
     lambda t: None if t == 0 else
     (f"={S('sub_share', t)}*{E('M_fuel', t)}*{F('ERavg', t)}" if t < 2 else
      f"=({S('sub_share', t)}+{P('sub_return')})*{E('M_fuel', t)}*{F('ERavg', t)}")),
    ("pb_nc", "Balance primario no-commodity", "% PIB", "pct",
     lambda t: f"={S('pb_nc', t)}" if t < 2 else f"={S('pb_nc', t)}+{P('fis_slip')}"),
    ("PB", "Balance primario", "Bs mm M", "bs",
     lambda t: f"={SC('pb0')}*{F('Y', t)}" if t == 0 else
     f"={F('pb_nc', t)}*{F('Y', t)}+{F('rev_h', t)}+{F('rev_m', t)}-{F('sub', t)}"),
    ("int_ext", "Intereses deuda externa", "Bs mm M", "bs",
     lambda t: None if t == 0 else f"={S('i_ext', t)}*{F('Dext', t - 1)}*{F('ERavg', t)}"),
    ("int_dom", "Intereses deuda interna", "Bs mm M", "bs",
     lambda t: None if t == 0 else f"={S('i_dom', t)}*{F('Ddom', t - 1)}"),
    ("OB", "Balance global", "Bs mm M", "bs",
     lambda t: f"={SC('ob0')}*{F('Y', t)}" if t == 0 else f"={F('PB', t)}-{F('int_ext', t)}-{F('int_dom', t)}"),
    ("netext", "Financiamiento externo neto", "Bs mm M", "bs",
     lambda t: None if t == 0 else f"=({E('disb', t)}-{E('amort', t)})*{F('ERavg', t)}"),
    ("cont", "Pasivos contingentes reconocidos", "Bs mm M", "bs",
     lambda t: None if t == 0 else (f"={P('cont')}*{F('Y', t)}" if t == 2 else 0.0)),
    ("domfin", "Financiamiento interno", "Bs mm M", "bs",
     lambda t: None if t == 0 else f"=-{F('OB', t)}-{F('netext', t)}+{F('cont', t)}"),
    ("Deuda pública SPNF", None),
    ("Dext", "Deuda externa pública", "USD mm M", "usd",
     lambda t: f"={SC('dext0')}" if t == 0 else f"={F('Dext', t - 1)}+{E('disb', t)}-{E('amort', t)}"),
    ("Ddom_usd", "Deuda interna en USD / indexada", "USD mm M", "usd", lambda t: f"={S('ddom_usd', t)}"),
    ("Ddom_lc", "Deuda interna en moneda local", "Bs mm M", "bs",
     lambda t: f"={SC('ddom_lc0')}" if t == 0 else f"={F('Ddom_lc', t - 1)}+{F('domfin', t)}"),
    ("Ddom", "Deuda interna total", "Bs mm M", "bs",
     lambda t: f"={F('Ddom_lc', t)}+{F('Ddom_usd', t)}*{F('ER', t)}"),
    ("D", "Deuda pública total", "Bs mm M", "bs",
     lambda t: f"={F('Dext', t)}*{F('ER', t)}+{F('Ddom', t)}"),
    ("d_gdp", "Deuda pública / PIB", "% PIB", "pct", lambda t: f"={F('D', t)}/{F('Y', t)}"),
    ("fx_share", "Deuda en moneda extranjera o indexada / total", "%", "pct",
     lambda t: f"=({F('Dext', t)}+{F('Ddom_usd', t)})*{F('ER', t)}/{F('D', t)}"),
    ("Dx_bcb", "Deuda neta de tenencias del BCB", "Bs mm M", "bs",
     lambda t: None if t == 0 else f"={F('D', t)}-{S('bcb_share', t)}*{F('Ddom_lc', t)}"),
    ("dx_gdp", "Deuda neta de BCB / PIB", "% PIB", "pct",
     lambda t: None if t == 0 else f"={F('Dx_bcb', t)}/{F('Y', t)}"),
    ("dom_amort", "Amortizaciones deuda interna", "Bs mm M", "bs",
     lambda t: None if t == 0 else f"={F('Ddom', t - 1)}/{SC('dom_mat')}"),
    ("GFN", "Necesidades brutas de financiamiento", "Bs mm M", "bs",
     lambda t: None if t == 0 else f"=-{F('OB', t)}+{F('dom_amort', t)}+{E('amort', t)}*{F('ERavg', t)}"),
    ("gfn_gdp", "NBF / PIB", "% PIB", "pct",
     lambda t: None if t == 0 else f"={F('GFN', t)}/{F('Y', t)}"),
    ("ob_gdp", "Balance global / PIB", "% PIB", "pct", lambda t: f"={F('OB', t)}/{F('Y', t)}"),
    ("pb_gdp", "Balance primario / PIB", "% PIB", "pct", lambda t: f"={F('PB', t)}/{F('Y', t)}"),
    ("Descomposición de la variación de deuda/PIB (pp)", None),
    ("c_pb", "Déficit primario", "pp", "pct",
     lambda t: None if t == 0 else f"=-{F('PB', t)}/{F('Y', t)}"),
    ("c_int", "Intereses", "pp", "pct",
     lambda t: None if t == 0 else f"=({F('int_ext', t)}+{F('int_dom', t)})/{F('Y', t)}"),
    ("c_g", "Crecimiento real", "pp", "pct",
     lambda t: None if t == 0 else
     f"=-{F('d_gdp', t - 1)}*{F('g', t)}/((1+{F('g', t)})*(1+{F('pi', t)}))"),
    ("c_pi", "Inflación (deflactor)", "pp", "pct",
     lambda t: None if t == 0 else
     f"=-{F('d_gdp', t - 1)}*{F('pi', t)}*(1+{F('g', t)})/((1+{F('g', t)})*(1+{F('pi', t)}))"),
    ("c_er", "Tipo de cambio (revaluación deuda en USD)", "pp", "pct",
     lambda t: None if t == 0 else
     (f"=(({F('Dext', t - 1)}+{F('Ddom_usd', t - 1)})*({F('ER', t)}-{F('ER', t - 1)})"
      f"+({E('disb', t)}-{E('amort', t)})*({F('ER', t)}-{F('ERavg', t)})"
      f"+({F('Ddom_usd', t)}-{F('Ddom_usd', t - 1)})*{F('ER', t)})/{F('Y', t)}")),
    ("c_cont", "Pasivos contingentes", "pp", "pct",
     lambda t: None if t == 0 else f"={F('cont', t)}/{F('Y', t)}"),
    ("c_res", "Residuo (control = 0)", "pp", "pct",
     lambda t: None if t == 0 else
     (f"=({F('d_gdp', t)}-{F('d_gdp', t - 1)})-({F('c_pb', t)}+{F('c_int', t)}+{F('c_g', t)}"
      f"+{F('c_pi', t)}+{F('c_er', t)}+{F('c_cont', t)})")),
    ("dd", "Variación deuda/PIB", "pp", "pct",
     lambda t: None if t == 0 else f"={F('d_gdp', t)}-{F('d_gdp', t - 1)}"),
]

IND_SPEC = [
    ("Solvencia y liquidez fiscal", None),
    ("i_d", "Deuda pública / PIB", "% PIB", "pct", lambda t: f"={F('d_gdp', t)}"),
    ("i_dx", "Deuda neta de tenencias BCB / PIB", "% PIB", "pct", lambda t: None if t == 0 else f"={F('dx_gdp', t)}"),
    ("i_gfn", "NBF / PIB", "% PIB", "pct", lambda t: None if t == 0 else f"={F('gfn_gdp', t)}"),
    ("i_ob", "Balance global / PIB", "% PIB", "pct", lambda t: f"={F('ob_gdp', t)}"),
    ("i_fx", "Deuda en ME o indexada / total", "%", "pct", lambda t: f"={F('fx_share', t)}"),
    ("Sostenibilidad externa", None),
    ("priv_ext", "Deuda externa privada", "USD mm M", "usd",
     lambda t: f"={SC('priv_ext0')}" if t < 2 else f"={I('priv_ext', t - 1)}*{F('Yusd', t)}/{F('Yusd', t - 1)}"),
    ("ED", "Deuda externa total (pública + FMI + privada)", "USD mm M", "usd",
     lambda t: f"={F('Dext', t)}+{E('imf_out', t)}+{I('priv_ext', t)}"),
    ("ed_gdp", "Deuda externa total / PIB", "% PIB", "pct", lambda t: f"={I('ED', t)}/{F('Yusd', t)}"),
    ("ed_x", "Deuda externa total / export. bienes", "%", "pct", lambda t: f"={I('ED', t)}/{E('X', t)}"),
    ("ds", "Servicio deuda externa pública + FMI", "USD mm M", "usd",
     lambda t: None if t == 0 else f"=-{E('pubint', t)}-{E('imfch', t)}+{E('amort', t)}+{E('imf_rep', t)}"),
    ("ds_x", "Servicio deuda externa / export. bienes", "%", "pct",
     lambda t: None if t == 0 else f"={I('ds', t)}/{E('X', t)}"),
    ("ca_gdp", "Cuenta corriente / PIB", "% PIB", "pct", lambda t: f"={E('CA', t)}/{F('Yusd', t)}"),
    ("res", "Reservas brutas", "USD mm M", "usd", lambda t: f"={E('R', t)}"),
    ("res_m", "Reservas / meses de importación", "meses", "m", lambda t: f"={E('R', t)}/({E('M', t)}/12)"),
    ("gg", "Reservas / vencimientos externos del año (Greenspan-Guidotti)", "x", "x",
     lambda t: None if t == 0 else
     f"={E('R', t)}/({E('amort', t)}+{E('imf_rep', t)}+{SC('st_share')}*{I('priv_ext', t)})"),
    ("Ajuste cambiario endógeno", None),
    ("i_gap", "Brecha externa cerrada vía tipo de cambio", "USD mm M", "usd", lambda t: f"={E('gap', t)}"),
    ("i_d_dep", "Depreciación de ajuste", "%", "pct", lambda t: f"={E('d', t)}"),
    ("i_er", "Tipo de cambio fin de año", "Bs/USD", "er", lambda t: f"={F('ER', t)}"),
]

# umbrales de referencia (valor, dirección: 'gt' = alarma si mayor)
THRESH = {
    "i_d": (0.70, "gt", "70%: umbral de alto escrutinio, marco MAC SRDSF (mercados emergentes)"),
    "i_gfn": (0.15, "gt", "15%: umbral NBF, marco MAC SRDSF"),
    "i_fx": (0.50, "gt", "Referencia: >50% eleva vulnerabilidad cambiaria"),
    "ed_gdp": (0.40, "gt", "40%: referencia LIC-DSF (capacidad media, en VP)"),
    "ed_x": (1.80, "gt", "180%: referencia LIC-DSF (capacidad media, en VP)"),
    "ds_x": (0.15, "gt", "15%: referencia LIC-DSF (capacidad media)"),
    "res_m": (3.0, "lt", "3 meses de importaciones"),
    "gg": (1.0, "lt", "Regla Greenspan-Guidotti: >1"),
}


def build_sheet(ws, sheet, spec, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT, italic=True, size=9)
    year_header(ws, 4)
    r = 5
    for item in spec:
        if item[1] is None:
            r += 1
            continue
        ROWS[sheet][item[0]] = r
        r += 1


def fill_sheet(ws, sheet, spec):
    CUR[0] = sheet
    r = 5
    for item in spec:
        if item[1] is None:
            section(ws, r, item[0])
            r += 1
            continue
        key, label, unit, fmt, fn = item
        note = THRESH[key][2] if key in THRESH else None
        write_row(ws, sheet, r, key, label, unit, fn, fmt, note,
                  bold=key in ("CA", "R", "d_gdp", "X", "gap", "OB", "i_d", "i_gfn", "res_m"))
        if key in THRESH:
            th, direction, _ = THRESH[key]
            rng = f"{col(0)}{r}:{col(T - 1)}{r}"
            op = "greaterThan" if direction == "gt" else "lessThan"
            ws.conditional_formatting.add(rng, CellIsRule(operator=op, formula=[str(th)], fill=RED_FILL))
        r += 1
    setw(ws, {"A": 56, "B": 12, **{col(t): 10 for t in range(T)}, col(T): 60})
    ws.freeze_panes = "C5"


build_sheet(ws_ext, "ext", EXT_SPEC, "Bloque externo: balanza de pagos, reservas y ajuste cambiario",
            "USD miles de millones salvo indicación. La brecha frente al nivel cómodo/piso de reservas se cierra con "
            "depreciación endógena, que retroalimenta deuda, crecimiento e inflación (hoja Fiscal_Deuda).")
build_sheet(ws_fis, "fis", FIS_SPEC, "Bloque fiscal y dinámica de deuda pública (SPNF)",
            "Bs miles de millones salvo indicación. Ingresos por hidrocarburos y minería, y subsidios a combustibles, "
            "dependen del bloque externo y del tipo de cambio.")
build_sheet(ws_ind, "ind", IND_SPEC, "Indicadores de sostenibilidad (escenario activo)",
            "En rojo: valores que superan umbrales de referencia (columna de notas).")
fill_sheet(ws_ext, "ext", EXT_SPEC)
fill_sheet(ws_fis, "fis", FIS_SPEC)
fill_sheet(ws_ind, "ind", IND_SPEC)

# ---------------------------------------------------------------------------
# PORTADA
# ---------------------------------------------------------------------------
ws = ws_port
ws["A1"] = "Bolivia: análisis de sostenibilidad de deuda con bloque externo integrado"
ws["A1"].font = Font(name=FONT, bold=True, size=16)
ws["A2"] = ("Modelo anual 2025-2035. Deuda pública del SPNF + balanza de pagos + reservas + tipo de cambio endógeno. "
            "Elaborado el 14-oct-2026 con información pública disponible a esa fecha.")
ws["A2"].font = Font(name=FONT, italic=True, size=9)
ws["A4"] = "Cómo usarlo"
ws["A4"].font = BOLD
ws["A6"] = "Escenario activo (1-8) ►"
ws["A6"].font = BOLD
ws["C6"] = ACTIVE
ws["C6"].font = Font(name=FONT, bold=True, color="0000FF", size=12)
ws["C6"].fill = KEY_FILL
ws["D6"] = f"=INDEX(Supuestos!$B${scen_first}:$B${scen_last},C6)"
ws["D6"].font = BOLD
dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(len(SCEN)), allow_blank=False)
ws.add_data_validation(dv)
dv.add("C6")
ws["A5"] = ("Cambie C6 para ver cualquier escenario en las hojas Externo, Fiscal_Deuda e Indicadores. "
            "Insumos en azul (hoja Supuestos). Las hojas Escenarios y Estocastico son salidas del motor en Python (modelo.py), "
            "con las mismas ecuaciones.")
ws["A5"].font = Font(name=FONT, size=9)
ws["A5"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A5:H5")
ws.row_dimensions[5].height = 30

ws["A8"] = "Resultados clave (escenario activo, fórmulas vivas)"
ws["A8"].font = BOLD
key_years = [2026, 2027, 2029, 2031, 2035]
ws.cell(row=9, column=1, value="Indicador").font = HDR_FONT
for j, y in enumerate(key_years):
    ws.cell(row=9, column=3 + j, value=str(y)).font = HDR_FONT
for cc in range(1, 3 + len(key_years)):
    ws.cell(row=9, column=cc).fill = HDR_FILL
KEY_IND = [("i_d", "Deuda pública / PIB", "pct"), ("i_dx", "Deuda neta de BCB / PIB", "pct"),
           ("i_gfn", "NBF / PIB", "pct"), ("i_ob", "Balance global / PIB", "pct"),
           ("ed_gdp", "Deuda externa total / PIB", "pct"), ("ds_x", "Servicio deuda ext. / export.", "pct"),
           ("ca_gdp", "Cuenta corriente / PIB", "pct"), ("res", "Reservas brutas (USD mm M)", "usd"),
           ("res_m", "Reservas (meses de import.)", "m"), ("i_er", "Tipo de cambio (Bs/USD)", "er"),
           ("i_d_dep", "Depreciación de ajuste endógena", "pct")]
for i, (k, lab, fmt) in enumerate(KEY_IND):
    rr = 10 + i
    ws.cell(row=rr, column=1, value=lab).font = BLACK
    for j, y in enumerate(key_years):
        t = YEARS.index(y)
        c = ws.cell(row=rr, column=3 + j, value=f"=Indicadores!{col(t)}{ROWS['ind'][k]}")
        c.font = GREEN
        c.number_format = FMT[fmt]

rr = 10 + len(KEY_IND) + 1
ws.cell(row=rr, column=1, value="Lectura de colores").font = BOLD
legend = [("Azul", "insumo editable", BLUE), ("Negro", "fórmula", BLACK), ("Verde", "vínculo a otra hoja", GREEN),
          ("Fondo amarillo", "supuesto clave / selector", BLACK), ("Fondo gris", "2025 dato, 2026 estimación", BLACK)]
for i, (a, b, f) in enumerate(legend):
    ws.cell(row=rr + 1 + i, column=1, value=a).font = f
    ws.cell(row=rr + 1 + i, column=3, value=b).font = BLACK
ws.cell(row=rr + 1 + 3, column=1).fill = KEY_FILL
ws.cell(row=rr + 1 + 4, column=1).fill = EST_FILL
setw(ws, {"A": 40, "B": 4, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12})

# ---------------------------------------------------------------------------
# ESCENARIOS (salida del motor Python)
# ---------------------------------------------------------------------------
results = [(name, m.run(q)) for name, q in SCEN]
ws = ws_esc
ws["A1"] = "Comparación de escenarios (salida de modelo.py)"
ws["A1"].font = TITLE
ws["A2"] = ("Valores calculados por el motor en Python con las mismas ecuaciones que las hojas de fórmulas; "
            "seleccionando el escenario en Portada!C6 las hojas vivas reproducen estas cifras.")
ws["A2"].font = Font(name=FONT, italic=True, size=9)
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "6250D6", "E34948"]
blocks = [("d_gdp", "Deuda pública / PIB", "pct"), ("R", "Reservas brutas (USD mm M)", "usd"),
          ("ER", "Tipo de cambio fin de año (Bs/USD)", "er"), ("gfn_gdp", "NBF / PIB", "pct"),
          ("ca_gdp", "Cuenta corriente / PIB", "pct"), ("ds_x", "Servicio deuda externa / export.", "pct"),
          ("d", "Depreciación de ajuste endógena", "pct")]
r = 4
block_rows = {}
for key, lab, fmt in blocks:
    ws.cell(row=r, column=1, value=lab).font = BOLD
    for cc in range(1, FIRST_COL + T):
        ws.cell(row=r, column=cc).fill = SEC_FILL
    r += 1
    ws.cell(row=r, column=1, value="Escenario").font = HDR_FONT
    for t, y in enumerate(YEARS):
        ws.cell(row=r, column=FIRST_COL + t, value=str(y)).font = HDR_FONT
    for cc in range(1, FIRST_COL + T):
        ws.cell(row=r, column=cc).fill = HDR_FILL
    hdr = r
    r += 1
    for name, o in results:
        ws.cell(row=r, column=1, value=name).font = BLACK
        for t in range(T):
            v = float(o[key][0, t])
            if key in ("ds_x", "gfn_gdp", "d") and t == 0:
                continue
            c = ws.cell(row=r, column=FIRST_COL + t, value=v)
            c.number_format = FMT[fmt]
            c.font = BLACK
        r += 1
    block_rows[key] = (hdr, r - 1)
    r += 1
setw(ws, {"A": 44, "B": 2, **{col(t): 9 for t in range(T)}})


def add_line_chart(ws, title, ytitle, hdr, first, last, anchor, fmt, colors, width=22, height=11):
    ch = LineChart()
    ch.title = title
    ch.y_axis.title = ytitle
    ch.y_axis.number_format = fmt
    ch.height = height
    ch.width = width
    data = Reference(ws, min_col=1, max_col=FIRST_COL + T - 1, min_row=first, max_row=last)
    ch.add_data(data, from_rows=True, titles_from_data=True)
    # quitar columna B vacía del rango: se usa from_rows con títulos en col A
    cats = Reference(ws, min_col=FIRST_COL, max_col=FIRST_COL + T - 1, min_row=hdr, max_row=hdr)
    ch.set_categories(cats)
    for i, s in enumerate(ch.series):
        s.graphicalProperties.line.solidFill = colors[i % len(colors)]
        s.graphicalProperties.line.width = 22000
        s.smooth = False
    ch.legend.position = "b"
    ws.add_chart(ch, anchor)


# Las series en filas incluyen la columna B (vacía); usamos rangos desde C con títulos aparte
def add_chart_rows(ws, title, ytitle, hdr, first, last, anchor, fmt, colors, labels_col=1, ymin=None):
    from openpyxl.chart import Series
    ch = LineChart()
    ch.title = title
    ch.y_axis.title = ytitle
    ch.y_axis.number_format = fmt
    if ymin is not None:
        ch.y_axis.scaling.min = ymin
    ch.height = 11
    ch.width = 22
    cats = Reference(ws, min_col=FIRST_COL, max_col=FIRST_COL + T - 1, min_row=hdr, max_row=hdr)
    for i, rr in enumerate(range(first, last + 1)):
        vals = Reference(ws, min_col=FIRST_COL, max_col=FIRST_COL + T - 1, min_row=rr, max_row=rr)
        s = Series(vals, title=ws.cell(row=rr, column=labels_col).value)
        s.graphicalProperties.line.solidFill = colors[i % len(colors)]
        s.graphicalProperties.line.width = 22000
        s.smooth = False
        ch.series.append(s)
    ch.set_categories(cats)
    ch.legend.position = "b"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, anchor)


h, l = block_rows["d_gdp"]
add_chart_rows(ws, "Deuda pública / PIB por escenario", "% PIB", h, h + 1, l, "O4", "0%", PALETTE, ymin=0.6)
h, l = block_rows["R"]
add_chart_rows(ws, "Reservas brutas por escenario (USD mm M)", "USD mm M", h, h + 1, l, "O27", "0.0", PALETTE)
h, l = block_rows["ER"]
add_chart_rows(ws, "Tipo de cambio por escenario (Bs/USD)", "Bs/USD", h, h + 1, l, "O50", "0.0", PALETTE)

# ---------------------------------------------------------------------------
# ESTOCÁSTICO
# ---------------------------------------------------------------------------
N_MC = 5000
mc = m.monte_carlo(N_MC)
ws = ws_mc
ws["A1"] = "Simulación estocástica (Monte Carlo, 5.000 sendas, escenario base)"
ws["A1"].font = TITLE
ws["A2"] = ("Choques correlacionados a precios de metales preciosos/base, petróleo y crecimiento (AR(1) en niveles), "
            "más choques a volumen de gas, precios agro, flujos privados, balance primario y tasa interna. "
            "A diferencia del fan chart estándar del FMI (VAR sobre variables fiscales históricas), aquí los choques "
            "externos pasan por la balanza de pagos, las reservas y el tipo de cambio antes de llegar a la deuda.")
ws["A2"].font = Font(name=FONT, italic=True, size=9)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:N2")
ws.row_dimensions[2].height = 42
PCTS = [10, 25, 50, 75, 90]
r = 4
mc_rows = {}
for key, lab, fmt in [("d_gdp", "Deuda pública / PIB", "pct"), ("R", "Reservas brutas (USD mm M)", "usd"),
                      ("ER", "Tipo de cambio (Bs/USD)", "er"), ("gfn_gdp", "NBF / PIB", "pct")]:
    ws.cell(row=r, column=1, value=lab).font = BOLD
    for cc in range(1, FIRST_COL + T):
        ws.cell(row=r, column=cc).fill = SEC_FILL
    r += 1
    ws.cell(row=r, column=1, value="Percentil").font = HDR_FONT
    for t, y in enumerate(YEARS):
        ws.cell(row=r, column=FIRST_COL + t, value=str(y)).font = HDR_FONT
    for cc in range(1, FIRST_COL + T):
        ws.cell(row=r, column=cc).fill = HDR_FILL
    hdr = r
    r += 1
    q = np.percentile(mc[key], PCTS, axis=0)
    for i, pc in enumerate(PCTS):
        ws.cell(row=r, column=1, value=f"P{pc}").font = BLACK
        for t in range(T):
            if key == "gfn_gdp" and t == 0:
                continue
            c = ws.cell(row=r, column=FIRST_COL + t, value=float(q[i, t]))
            c.number_format = FMT[fmt]
            c.font = BLACK
        r += 1
    mc_rows[key] = (hdr, r - 1)
    r += 1

ws.cell(row=r, column=1, value="Probabilidades (sobre 5.000 sendas)").font = BOLD
for cc in range(1, FIRST_COL + T):
    ws.cell(row=r, column=cc).fill = SEC_FILL
r += 1
i26, i29, i30, i35 = (YEARS.index(y) for y in (2026, 2029, 2030, 2035))
probs = [
    ("Deuda/PIB 2030 > 100%", (mc["d_gdp"][:, i30] > 1.0).mean()),
    ("Deuda/PIB 2030 > deuda/PIB 2026 (no baja)", (mc["d_gdp"][:, i30] > mc["d_gdp"][:, i26]).mean()),
    ("Deuda/PIB 2035 > 70% (umbral MAC)", (mc["d_gdp"][:, i35] > 0.70).mean()),
    ("NBF/PIB > 15% en 2030", (mc["gfn_gdp"][:, i30] > 0.15).mean()),
    ("Reservas < 3 meses de import. algún año 2027-2031",
     (mc["res_m"][:, 2:7].min(axis=1) < 3.0).mean()),
    ("Depreciación de ajuste > 10% algún año 2027-2031", (mc["d"][:, 2:7].max(axis=1) > 0.10).mean()),
    ("Deuda/PIB 2029 > 89,2% (proyección FMI)", (mc["d_gdp"][:, i29] > 0.892).mean()),
]
for lab, v in probs:
    ws.cell(row=r, column=1, value=lab).font = BLACK
    c = ws.cell(row=r, column=FIRST_COL, value=float(v))
    c.number_format = "0%"
    c.font = BLACK
    r += 1
FAN = ["9EC5F4", "5A9BE8", "1F3864", "5A9BE8", "9EC5F4"]
h, l = mc_rows["d_gdp"]
add_chart_rows(ws, "Abanico: deuda pública / PIB", "% PIB", h, h + 1, l, "O4", "0%", FAN, ymin=0.6)
h, l = mc_rows["R"]
add_chart_rows(ws, "Abanico: reservas brutas (USD mm M)", "USD mm M", h, h + 1, l, "O27", "0.0", FAN)
setw(ws, {"A": 52, "B": 2, **{col(t): 9 for t in range(T)}})

# ---------------------------------------------------------------------------
# FMI vs MODELO
# ---------------------------------------------------------------------------
ws = ws_cmp
ws["A1"] = "Qué agrega este DSA frente al marco estándar del FMI"
ws["A1"].font = TITLE
rows = [
    ("Dimensión", "DSA estándar (MAC SRDSF / staff report)", "Este modelo"),
    ("Tipo de cambio", "Senda exógena; choque cambiario genérico como prueba de tensión.",
     "Endógeno: la brecha de balanza de pagos frente al piso/nivel cómodo de reservas se cierra con depreciación, "
     "que revalúa la deuda en USD, sube el deflactor y resta crecimiento."),
    ("Exportaciones", "Agregadas; choques de commodities como escenario puntual.",
     "Desagregadas por motor (metales preciosos, metales base, gas, agro, resto) con precio y volumen propios; "
     "gas con declinación de campos e indexación rezagada al petróleo."),
    ("Vínculo fiscal-externo", "Implícito.",
     "Explícito: renta hidrocarburífera y minera dependen del valor exportado y del tipo de cambio; "
     "el subsidio a combustibles depende del precio del petróleo y del tipo de cambio; "
     "el desliz fiscal aumenta importaciones (déficits gemelos)."),
    ("Reservas", "Se reportan; no retroalimentan la deuda.",
     "Variable de estado: acumulan cuenta corriente + IED + financiamiento oficial + FMI + flujos privados "
     "+ valuación del oro. Determinan el ajuste cambiario."),
    ("Abanico estocástico", "Choques históricos a crecimiento, tasas, tipo de cambio y balance primario.",
     "Choques a precios de commodities correlacionados con persistencia, que atraviesan la balanza de pagos "
     "antes de llegar a la deuda."),
    ("Perímetro de deuda", "SPNF bruto.",
     "SPNF bruto y neto de tenencias del BCB (≈55% de la deuda interna), más composición por moneda."),
    ("Indicadores externos", "Separados en DSA externo.",
     "Integrados: deuda externa total/PIB y /export., servicio/export., reservas en meses de importación, "
     "Greenspan-Guidotti."),
    ("Supuesto implícito de TCR", "No se explicita.",
     "Escenario 'Réplica FMI' calibra cuánta apreciación real 2027-29 se necesita para que, con los déficits del "
     "programa, la deuda llegue a 89,2% del PIB en 2029. Ese mismo escenario muestra el costo externo de esa apreciación."),
]
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = ws.cell(row=3 + i, column=1 + j, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = HDR_FONT if i == 0 else (BOLD if j == 0 else BLACK)
        if i == 0:
            c.fill = HDR_FILL
r = 3 + len(rows) + 1
ws.cell(row=r, column=1, value="Limitaciones").font = BOLD
lims = [
    "Datos de balanza de pagos 2026 son estimaciones a partir de comercio exterior ene-ago (INE) y del WEO oct-2026; no hay BoP oficial 2026.",
    "El cronograma de amortizaciones externas es una aproximación; no se tuvo acceso al informe de deuda externa del BCB ni al staff report completo del FMI.",
    "La regla cambiaria es asimétrica (sólo deprecia ante caída de reservas); con reservas abundantes el modelo acumula en vez de apreciar.",
    "No modela explícitamente el sistema financiero ni la dolarización de depósitos; los pasivos contingentes entran como choque de stock.",
    "El auge exportador de 2026 mezcla precios (plata, oro) y formalización de exportaciones antes informales tras la flotación; separar ambos es clave y no está resuelto.",
    "Las elasticidades son supuestos calibrados, no estimaciones econométricas.",
]
for i, t in enumerate(lims):
    c = ws.cell(row=r + 1 + i, column=1, value=f"• {t}")
    c.font = BLACK
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r + 1 + i, start_column=1, end_row=r + 1 + i, end_column=3)
    ws.row_dimensions[r + 1 + i].height = 30
setw(ws, {"A": 28, "B": 50, "C": 80})

# ---------------------------------------------------------------------------
# FUENTES
# ---------------------------------------------------------------------------
ws = ws_src
ws["A1"] = "Fuentes"
ws["A1"].font = TITLE
SRC = [
    ("FMI, acuerdo a nivel técnico SAF (29-jul-2026)", "https://www.imf.org/en/news/articles/2026/07/29/pr26268-bolivia-imf-reaches-staff-level-agreement-on-an-extended-fund-facility-arrangement"),
    ("Aprobación del Directorio FMI, 2-oct-2026 (USD 1.900 M; USD 214 M inmediatos)", "https://www.riotimesonline.com/bolivia-imf-loan-approved-214-million-expats-2026/"),
    ("Informe técnico FMI: deuda 100,2% PIB 2026; déficit 11,4%→9,2%→6,7%", "https://fmbolivia.com.bo/nacional/deuda-publica-de-bolivia-llegaria-al-100-2-del-pib-en-2026-segun-informe-tecnico-del-fmi-21327"),
    ("Acuerdo FMI: ajuste 8,5% PIB, deuda 89,2% en 2029", "https://fmlapaz.bo/fmi-hace-publico-el-acuerdo-con-bolivia-us-1-900-millones-ajuste-del-85-del-pib-y-eliminacion-del-subsidio-a-los-combustibles-en-2027/"),
    ("FMI WEO oct-2026: PIB -3,0%/-1,5%; CC +1,6%/+1,0% PIB", "https://www.vision360.bo/noticias/2026/10/13/72357-fmi-bolivia-sera-la-unica-economia-de-sudamerica-que-se-contraera-en-2026-y-2027"),
    ("Deuda externa pública USD 14.117 M a sep-2026 (BCB)", "https://unitel.bo/noticias/agencias/deuda-externa-publica-de-bolivia-llega-a-14117-millones-de-dolares-a-septiembre-de-2026-OC24040790"),
    ("BCB, Reporte de Balanza de Pagos 2025", "https://www.bcb.gob.bo/webdocs/publicacionesbcb/2026/03/31/BOP%202025%20(4).pdf"),
    ("BCB, Resumen de cifras BoP 2025 (CC -1.218 M; bienes +334 M)", "https://www.bcb.gob.bo/webdocs/publicacionesbcb/2026/02/03/Resumen%20de%20cifras%20(WEB).pdf"),
    ("INE vía eju.tv: export. ene-ago 2026 USD 8.987 M; saldo comercial +2.029 M", "https://eju.tv/2026/10/bolivia-registro-un-saldo-comercial-superior-a-us-2-000-millones-entre-enero-y-agosto-de-este-ano-segun-el-ine/"),
    ("IBCE: exportaciones a junio 2026", "https://www.ibce.org.bo//images/ibcecifras_documentos/Cifras_1.441_Exportaciones_junio_2026_1786634430.pdf"),
    ("Export. de gas 2025: USD 1.077 M (-33%)", "https://unitel.bo/noticias/economia/bolivia-registro-una-caida-cercana-a-us-540-millones-en-la-exportacion-de-gas-en-2025-segun-el-ine-PL19126328"),
    ("CEPAL: PIB nominal 2025 USD 55.000 M", "https://correodelsur.com/economia/20260227/cepal-el-pib-nominal-de-bolivia-llego-a-us-55-000-millones-en-2025.html"),
    ("Flotación cambiaria: RM 245/2026 y RD 88/2026", "https://www.dentons.com/en/insights/alerts/2026/july/1/bolivia-adopts-a-flexible-exchange-rate-regime"),
    ("BCB, tipos de cambio históricos 2026", "https://www.bcb.gob.bo/tiposDeCambioHistorico/pdf.php?anio=2026"),
    ("Reservas líquidas >USD 1.000 M tras desembolso FMI", "https://eju.tv/2026/10/reservas-liquidas-del-bcb-crecen-a-mas-de-us-1-000-millones-tras-el-desembolso-del-fmi/"),
    ("BCB mayor acreedor de deuda interna (54,6%)", "https://www.jubileobolivia.org.bo/Publicaciones/Presupuestos-y-Fiscalidad/Deuda-Interna:-El-BCB-es-el-mayor-acreedor-con-546"),
    ("Deuda interna TGN (MEFP)", "https://www.economiayfinanzas.gob.bo/viceministerios/vtcp/deuda-interna-tgn"),
    ("Canje de bonos 2028/2030 (mar-2026)", "https://riosmauricio.com/2026/04/algunas-medias-verdades-sobre-los-bonos-soberanos-2028-de-bolivia/"),
    ("Infobae: desembolsos esperados FMI/BM/BID", "https://www.infobae.com/america/america-latina/2026/09/19/el-congreso-de-bolivia-aprobo-el-acuerdo-de-financiamiento-con-el-fmi-por-usd-1900-millones/"),
    ("FinDev Lab (abr-2026): DSA independiente con reperfilamiento", "http://findevlab.org/wp-content/uploads/2026/04/FDL_Policy-Note-34_A-New-Day-for-Bolivia_The-Anatomy-of-a-Crisis-and-the-Options-Ahead.pdf"),
]
ws.cell(row=3, column=1, value="Dato").font = HDR_FONT
ws.cell(row=3, column=2, value="URL").font = HDR_FONT
ws.cell(row=3, column=1).fill = HDR_FILL
ws.cell(row=3, column=2).fill = HDR_FILL
for i, (a, b) in enumerate(SRC):
    ws.cell(row=4 + i, column=1, value=a).font = BLACK
    c = ws.cell(row=4 + i, column=2, value=b)
    c.hyperlink = b
    c.font = Font(name=FONT, color="0563C1", underline="single", size=9)
ws.cell(row=5 + len(SRC), column=1, value=(
    "Nota: varias cifras provienen de prensa que cita fuentes oficiales; el acceso directo a imf.org y bcb.gob.bo "
    "no estuvo disponible al construir el modelo. Verificar contra los documentos originales antes de publicar.")).font = Font(name=FONT, italic=True, size=9)
setw(ws, {"A": 70, "B": 110})

wb.save(OUT)
print("ok", OUT, "filas ext", len(ROWS["ext"]), "fis", len(ROWS["fis"]))
